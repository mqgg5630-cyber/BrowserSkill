#!/usr/bin/env python3
"""Generator for BrowserSkill self-loop deliverables (docx + pptx).

Usage:
  python3 scripts/make_deliverables.py --round 1 --prev-status "..." [--timestamp "..."]
"""
import argparse
import datetime
import os
import subprocess
import sys

def repo_info():
    def run(*args):
        try:
            return subprocess.check_output(args, text=True, cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))).strip()
        except Exception:
            return "n/a"
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    branch = run("git", "rev-parse", "--abbrev-ref", "HEAD")
    commit = run("git", "rev-parse", "--short", "HEAD")
    return root, branch, commit

def build_docx(path, round_no, total_rounds, timestamp, prev_status, branch, commit):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading(f"BrowserSkill 自循环验收报告 — 第 {round_no} 轮", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"分支 arena/01a0b3c5-browserskill · 提交 {commit} · {timestamp}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_heading("1. 本轮概览", level=2)
    rows = [
        ("本轮轮次", f"第 {round_no} 轮 / 共 {total_rounds} 轮"),
        ("生成时间戳 (UTC)", timestamp),
        ("上轮状态", prev_status),
        ("分支", "arena/01a0b3c5-browserskill"),
        ("基线提交", commit),
        ("验收方式", "push 后经 agent-handsfree 请求本机验收；passed + --accept 收尾"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "内容"
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = v

    doc.add_heading("2. 交付物清单 (本轮)", level=2)
    doc.add_paragraph(f"deliverables/round-{round_no}/report-round-{round_no}.docx（本文件）", style="List Bullet")
    doc.add_paragraph(f"deliverables/round-{round_no}/report-round-{round_no}.pptx（配套演示文稿）", style="List Bullet")
    doc.add_paragraph(f"deliverables/round-{round_no}/meta.json（轮次 / 时间戳 / 上轮状态 / 提交元数据）", style="List Bullet")
    doc.add_paragraph("scripts/agent-handsfree（验收请求与 --accept 收尾工具）", style="List Bullet")
    doc.add_paragraph("scripts/make_deliverables.py（交付物生成器）", style="List Bullet")

    doc.add_heading("3. 上轮状态说明", level=2)
    doc.add_paragraph(prev_status)
    if round_no == 1:
        doc.add_paragraph("第一轮自循环承接用户确认的准备态：用户本机已跑完第一轮命令块，bsk doctor 全绿。本轮为自循环第 1 轮，目标是建立可重复的交付 + 验收闭环。", style="List Bullet")
    else:
        doc.add_paragraph(f"上一轮已通过本机验收（passed）并执行 --accept 收尾，本轮在此基础上递增生成，不回退、不停顿。", style="List Bullet")

    doc.add_heading("4. 本轮验收标准", level=2)
    checks = [
        "docx / pptx 文件存在且非空，可被 Office / WPS 正常打开",
        "docx / pptx 内容包含：该轮轮次、时间戳、上轮状态三要素",
        "已 push 到分支 arena/01a0b3c5-browserskill，远端可拉取",
        "agent-handsfree request 返回 passed",
        "agent-handsfree --accept 已收尾，ACCEPTED.json 已生成并 push",
    ]
    for c in checks:
        doc.add_paragraph(c, style="List Number")

    doc.add_heading("5. 判定与下一步", level=2)
    if round_no < total_rounds:
        doc.add_paragraph(f"本轮通过后立即自动开始第 {round_no + 1} 轮，不停顿、不提问。3 轮全部通过后输出简短总结再停。")
    else:
        doc.add_paragraph("本轮为最后一轮（第 3 轮）。通过后输出三轮判定 + 交付物清单总结并停止。")
    doc.add_paragraph(f"生成工具: scripts/make_deliverables.py · 验收工具: scripts/agent-handsfree · 分支: {branch}")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path

def build_pptx(path, round_no, total_rounds, timestamp, prev_status, branch, commit):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_slide(title, bullets, note=None):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.level = 0
            for run in p.runs:
                run.font.size = Pt(18)
        if note:
            slide.notes_slide.placeholders[1].text = note
        return slide

    # Cover
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = f"BrowserSkill 自循环验收 — 第 {round_no} 轮"
    slide.placeholders[1].text = f"共 {total_rounds} 轮 · {timestamp}\n分支 arena/01a0b3c5-browserskill · {commit}"

    add_slide("本轮概览", [
        f"轮次：第 {round_no} 轮 / 共 {total_rounds} 轮",
        f"时间戳 (UTC)：{timestamp}",
        f"上轮状态：{prev_status}",
        f"分支：arena/01a0b3c5-browserskill（{branch}@{commit}）",
    ])
    add_slide("交付物清单", [
        f"deliverables/round-{round_no}/report-round-{round_no}.docx",
        f"deliverables/round-{round_no}/report-round-{round_no}.pptx（本文件）",
        f"deliverables/round-{round_no}/meta.json",
        "scripts/agent-handsfree + scripts/make_deliverables.py",
    ])
    add_slide("验收标准", [
        "docx / pptx 存在、非空、可打开",
        "内容含轮次 + 时间戳 + 上轮状态",
        "已 push 到分支，远端可拉取",
        "agent-handsfree request → passed",
        "agent-handsfree --accept → ACCEPTED.json",
    ])
    if round_no < total_rounds:
        add_slide("下一步", [
            f"本轮 passed + --accept 后立即开始第 {round_no + 1} 轮",
            "全程不提问、自主完成",
            "3 轮通过后输出总结再停",
        ])
    else:
        add_slide("下一步", [
            "本轮为最后一轮（第 3 轮）",
            "通过后输出三轮判定 + 交付物清单总结",
            "总结输出后停止",
        ])

    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    return path

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--round", dest="round_no", type=int, required=True)
    ap.add_argument("--prev-status", required=True)
    ap.add_argument("--timestamp", default=None)
    ap.add_argument("--total", type=int, default=3)
    args = ap.parse_args()
    if args.round_no < 1 or args.round_no > args.total:
        print(f"round must be 1..{args.total}", file=sys.stderr)
        sys.exit(2)
    ts = args.timestamp or (datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"))
    root, branch, commit = repo_info()
    outdir = os.path.join(root, "deliverables", f"round-{args.round_no}")
    os.makedirs(outdir, exist_ok=True)
    docx_path = os.path.join(outdir, f"report-round-{args.round_no}.docx")
    pptx_path = os.path.join(outdir, f"report-round-{args.round_no}.pptx")
    build_docx(docx_path, args.round_no, args.total, ts, args.prev_status, branch, commit)
    build_pptx(pptx_path, args.round_no, args.total, ts, args.prev_status, branch, commit)
    import json
    meta = {
        "round": args.round_no,
        "total_rounds": args.total,
        "timestamp_utc": ts,
        "prev_status": args.prev_status,
        "branch": "arena/01a0b3c5-browserskill",
        "commit_at_build": commit,
        "artifacts": [os.path.relpath(docx_path, root), os.path.relpath(pptx_path, root)],
    }
    with open(os.path.join(outdir, "meta.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    print(f"round={args.round_no} timestamp={ts}")
    print(f"docx={docx_path} ({os.path.getsize(docx_path)} bytes)")
    print(f"pptx={pptx_path} ({os.path.getsize(pptx_path)} bytes)")

if __name__ == "__main__":
    main()
